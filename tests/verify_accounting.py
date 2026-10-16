"""Hisob-kitobni tekshiradi. Yozuvlarni commit qilmaydi."""
import asyncio
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.database import Expense, Order, Payment, SessionLocal, TeamDebt, TeamMember, TeamPayment, engine
from app.excel_report import build_workbook
from app.formatting import (
    as_moment,
    month_bounds,
    parse_amount,
    parse_date,
    previous_month,
    status_for,
    today_local,
)
from app.reminders import due_event_orders, due_orders
from app.order_pdf import ClientPage, write_client_pdf, write_orders_pdf
from app.reports import (
    OrderLine,
    daily_period,
    debtors,
    excel_data,
    figures_for,
    monthly_clients,
    monthly_period,
)
from app.services import (
    UserFacingError,
    add_expense,
    add_payment,
    add_team_payment,
    create_order,
    create_team_debt,
    debt_remaining,
    delete_order,
    find_by_key,
    get_or_create_member,
    paid_sum,
    pay_team_debt,
    record_member_cash,
    settle_order,
    set_payment_due,
    set_team_debt_date,
)

TODAY = today_local()
PREVIOUS_YEAR, PREVIOUS_MONTH = previous_month(TODAY)
PREVIOUS_DAY = date(PREVIOUS_YEAR, PREVIOUS_MONTH, 15)


async def counts() -> tuple[int, int, int, int, int]:
    async with SessionLocal() as session:
        orders = int(await session.scalar(select(func.count(Order.id))) or 0)
        payments = int(await session.scalar(select(func.count(Payment.id))) or 0)
        expenses = int(await session.scalar(select(func.count(Expense.id))) or 0)
        team = int(await session.scalar(select(func.count(TeamPayment.id))) or 0)
        members = int(await session.scalar(select(func.count(TeamMember.id))) or 0)
    return orders, payments, expenses, team, members


def check_money() -> None:
    assert parse_amount("-5") is None
    assert parse_amount("0") is None
    assert parse_amount("0", allow_zero=True) == Decimal("0.00")
    assert parse_amount("1,5") == Decimal("1.50")
    assert parse_amount("1.555") == Decimal("1555.00")
    assert parse_amount("1.5555") is None
    assert parse_amount("1500000") == Decimal("1500000.00")
    same = Decimal("100000.00")
    assert parse_amount("100.000") == same
    assert parse_amount("100 000") == same
    assert parse_amount("100000") == same
    assert parse_amount("100,000") == same
    assert parse_amount("100.000 so'm") == same
    assert parse_amount("1.500.000") == Decimal("1500000.00")
    assert parse_amount("1500.50") == Decimal("1500.50")
    assert parse_amount("1500,50") == Decimal("1500.50")
    assert status_for(100, 0) == "Yangi"
    assert status_for(100, 40) == "Qisman to'langan"
    assert status_for(100, 100) == "Yopilgan"
    start, end = month_bounds(2026, 10)
    assert start == date(2026, 10, 1)
    assert end == date(2026, 11, 1)
    same_day = date(2026, 12, 12)
    assert parse_date("12-12-2026") == same_day
    assert parse_date("12.12.2026") == same_day
    assert parse_date("12 12 2026") == same_day
    assert parse_date("12/12/2026") == same_day
    assert parse_date("32-13-2026") is None


async def check_empty_excel() -> None:
    async with SessionLocal() as session:
        data = await excel_data(session, daily_period(TODAY))
    workbook = build_workbook(data)
    names = workbook.sheetnames
    assert names == [
        "Umumiy hisobot",
        "Buyurtmalar",
        "To'lovlar",
        "Qarzdorlar",
        "Xarajatlar",
        "Jamoa to'lovlari",
    ]
    sheet = workbook["Buyurtmalar"]
    assert sheet["A1"].value == "FAZO CREATIVE GROUP"
    assert sheet.freeze_panes == "A5"
    assert sheet.auto_filter.ref is not None
    assert any(sheet.cell(row, 1).value == "Jami" for row in range(5, sheet.max_row + 1))
    workbook.close()


async def check_accounting() -> None:
    async with SessionLocal() as session:
        try:
            daily_before = await figures_for(session, daily_period(TODAY))
            month = monthly_period(TODAY.year, TODAY.month)
            monthly_before = await figures_for(session, month)
            order, created = await create_order(
                session,
                client_name="VERIFY FAZO",
                phone="+998901112233",
                service="Foto xizmat",
                total_price=Decimal("1000000.00"),
                prepaid=Decimal("200000.00"),
                event_date=TODAY,
                notes="tekshiruv",
                idempotency_key="verify-order-1",
            )
            assert created is True
            assert order.order_number == f"FCG-{order.id:05d}"
            assert await paid_sum(session, order.id) == Decimal("200000.00")

            again, created_again = await create_order(
                session,
                client_name="VERIFY FAZO",
                phone="+998901112233",
                service="Foto xizmat",
                total_price=Decimal("1000000.00"),
                prepaid=Decimal("200000.00"),
                event_date=TODAY,
                notes="tekshiruv",
                idempotency_key="verify-order-1",
            )
            assert created_again is False
            assert again.id == order.id

            await add_payment(
                session,
                order_id=order.id,
                amount=Decimal("300000.00"),
                method="Naqd",
                note="bugun",
                paid_at=as_moment(TODAY),
                idempotency_key="verify-pay-today",
            )
            duplicate, _, duplicate_created = await add_payment(
                session,
                order_id=order.id,
                amount=Decimal("300000.00"),
                method="Naqd",
                note="bugun",
                paid_at=as_moment(TODAY),
                idempotency_key="verify-pay-today",
            )
            assert duplicate_created is False
            assert duplicate.amount == Decimal("300000.00")

            await add_payment(
                session,
                order_id=order.id,
                amount=Decimal("100000.00"),
                method="Karta",
                note="oldingi oy",
                paid_at=as_moment(PREVIOUS_DAY),
                idempotency_key="verify-pay-prev",
            )
            assert await paid_sum(session, order.id) == Decimal("600000.00")
            try:
                await add_payment(
                    session,
                    order_id=order.id,
                    amount=Decimal("500000.00"),
                    method="Naqd",
                    note="ortiqcha",
                    paid_at=as_moment(TODAY),
                    idempotency_key="verify-pay-over",
                )
            except UserFacingError:
                pass
            else:
                raise AssertionError("Ortiqcha to'lov rad etilmadi")
            assert await paid_sum(session, order.id) == Decimal("600000.00")

            closed, _ = await create_order(
                session,
                client_name="VERIFY YOPILGAN",
                phone="+998901110000",
                service="Video xizmat",
                total_price=Decimal("50000.00"),
                prepaid=Decimal("50000.00"),
                event_date=None,
                notes="",
                idempotency_key="verify-order-closed",
            )
            assert closed.status == "Yopilgan"

            await add_expense(
                session,
                category="Transport",
                amount=Decimal("100000.00"),
                description="yo'l",
                spent_at=as_moment(TODAY),
                idempotency_key="verify-expense",
            )
            await add_team_payment(
                session,
                person_name="VERIFY XODIM",
                amount=Decimal("50000.00"),
                work_for=order.order_number or "",
                order_id=order.id,
                note="smena",
                paid_at=as_moment(TODAY),
                idempotency_key="verify-team",
            )

            daily = await figures_for(session, daily_period(TODAY))
            assert daily.new_orders == daily_before.new_orders + 2
            assert daily.order_value == daily_before.order_value + Decimal("1050000.00")
            assert daily.received == daily_before.received + Decimal("550000.00")
            assert daily.expenses == daily_before.expenses + Decimal("100000.00")
            assert daily.team_paid == daily_before.team_paid + Decimal("50000.00")
            assert daily.outgoing == daily_before.outgoing + Decimal("150000.00")
            assert daily.net == daily_before.net + Decimal("400000.00")
            assert daily.outstanding_debt == daily_before.outstanding_debt + Decimal("400000.00")
            assert daily.debtor_count == daily_before.debtor_count + 1

            monthly = await figures_for(session, month)
            assert monthly.received == monthly_before.received + Decimal("550000.00")
            clients = await monthly_clients(session, TODAY.year, TODAY.month)
            debtor = next(item for item in clients if item.client_name == "VERIFY FAZO")
            assert debtor.paid_in_period == Decimal("500000.00")
            assert debtor.paid_before == Decimal("100000.00")
            assert debtor.paid_total == Decimal("600000.00")
            assert debtor.debt == Decimal("400000.00")
            assert debtor.total_price == Decimal("1000000.00")

            debtor_lines = await debtors(session)
            names = {line.client_name for line in debtor_lines}
            assert "VERIFY FAZO" in names
            assert "VERIFY YOPILGAN" not in names

            workbook = build_workbook(await excel_data(session, month))
            team_sheet = workbook["Jamoa to'lovlari"]
            expense_sheet = workbook["Xarajatlar"]
            team_rows = [
                row
                for row in team_sheet.iter_rows(min_row=5, max_col=3, values_only=True)
                if row[1] == "VERIFY XODIM"
            ]
            assert team_rows and Decimal(str(team_rows[0][2])) == Decimal("50000.00")
            assert "SUBTOTAL" in str(team_sheet.cell(team_sheet.max_row, 3).value)
            assert "SUBTOTAL" in str(expense_sheet.cell(expense_sheet.max_row, 3).value)
            workbook.close()

            member = await get_or_create_member(session, "VERIFY QARZ")
            debt, debt_created = await create_team_debt(
                session,
                member_id=member.id,
                amount=Decimal("80000.00"),
                order_id=order.id,
                work_for=order.order_number or "",
                opened_on=TODAY,
                idempotency_key="verify-debt",
            )
            assert debt_created is True
            _, again_created = await create_team_debt(
                session,
                member_id=member.id,
                amount=Decimal("80000.00"),
                order_id=order.id,
                work_for=order.order_number or "",
                opened_on=TODAY,
                idempotency_key="verify-debt",
            )
            assert again_created is False
            _, updated, pay_created = await pay_team_debt(
                session,
                debt_id=debt.id,
                amount=Decimal("30000.00"),
                paid_at=as_moment(TODAY),
                idempotency_key="verify-debt-pay",
            )
            assert pay_created is True
            assert debt_remaining(updated) == Decimal("50000.00")
            try:
                await pay_team_debt(
                    session,
                    debt_id=debt.id,
                    amount=Decimal("60000.00"),
                    paid_at=as_moment(TODAY),
                    idempotency_key="verify-debt-over",
                )
            except UserFacingError:
                pass
            else:
                raise AssertionError("Ortish qarz rad etilmadi")
            assert debt_remaining(updated) == Decimal("50000.00")
            await set_team_debt_date(session, debt.id, TODAY)
            assert debt.opened_on == TODAY
            member.phone = "+998901110099"
            assert member.phone == "+998901110099"
            await set_payment_due(session, order.id, TODAY)
            due = await due_orders(session)
            assert any(item.id == order.id for item in due)
            order.reminder_sent_on = TODAY
            due_after = await due_orders(session)
            assert all(item.id != order.id for item in due_after)

            removable, _ = await create_order(
                session,
                client_name="VERIFY OCHIRISH",
                phone="+998901110077",
                service="Foto xizmat",
                total_price=Decimal("10000.00"),
                prepaid=Decimal("0.00"),
                event_date=None,
                notes="",
                idempotency_key="verify-remove",
            )
            _, settled, settled_created = await settle_order(
                session,
                order_id=removable.id,
                paid_at=as_moment(TODAY),
                idempotency_key="verify-settle",
            )
            assert settled_created is True
            assert settled.status == "Yopilgan"
            assert await paid_sum(session, removable.id) == Decimal("10000.00")
            await add_team_payment(
                session,
                person_name="VERIFY XODIM",
                amount=Decimal("1000.00"),
                work_for=removable.order_number or "",
                order_id=removable.id,
                note="ochirish",
                paid_at=as_moment(TODAY),
                idempotency_key="verify-remove-team",
            )
            await create_team_debt(
                session,
                member_id=member.id,
                amount=Decimal("1000.00"),
                order_id=removable.id,
                work_for=removable.order_number or "",
                opened_on=TODAY,
                idempotency_key="verify-remove-debt",
            )
            removed_name = await delete_order(session, removable.id)
            assert removed_name == "VERIFY OCHIRISH"
            assert await session.get(Order, removable.id) is None
            assert await paid_sum(session, removable.id) == Decimal("0.00")
            kept_payment = await find_by_key(session, TeamPayment, "verify-remove-team")
            kept_debt = await find_by_key(session, TeamDebt, "verify-remove-debt")
            assert kept_payment is not None and kept_payment.order_id is None
            assert kept_debt is not None and kept_debt.order_id is None

            service_day, _ = await create_order(
                session,
                client_name="VERIFY TOY",
                phone="+998901110088",
                service="Foto xizmat",
                total_price=Decimal("20000.00"),
                prepaid=Decimal("0.00"),
                event_date=TODAY,
                notes="",
                idempotency_key="verify-service-day",
            )
            asked = await due_event_orders(session)
            assert any(item.id == service_day.id for item in asked)
            service_day.event_checked_on = TODAY
            asked_after = await due_event_orders(session)
            assert all(item.id != service_day.id for item in asked_after)
            written = await record_member_cash(
                session,
                member_id=member.id,
                order_id=service_day.id,
                amount=Decimal("7000.00"),
                paid_at=as_moment(TODAY),
                idempotency_key="verify-cash",
            )
            assert written is True
            cash = await find_by_key(session, TeamPayment, "verify-cash:cash")
            assert cash is not None
            assert cash.order_id == service_day.id
            assert cash.amount == Decimal("7000.00")
        finally:
            await session.rollback()


def check_orders_pdf() -> None:
    line = OrderLine(
        number="FCG-00001",
        client_name="Ali Valiyev",
        phone="+998901112233",
        service="Foto xizmati — 100 000 so'm, Video xizmati — 200 000 so'm",
        total_price=Decimal("300000.00"),
        paid=Decimal("100000.00"),
        debt=Decimal("200000.00"),
        event_date=None,
        created_on=None,
        status="Qisman to'langan",
        notes="",
    )
    path = write_orders_pdf([line])
    try:
        raw = path.read_bytes()
        assert raw.startswith(b"%PDF")
        assert b"FAZO CREATIVE GROUP" in raw
        assert len(raw) > 1000
    finally:
        path.unlink(missing_ok=True)
    client_path = write_client_pdf(
        ClientPage(
            number=line.number,
            client_name=line.client_name,
            phone=line.phone,
            service=line.service,
            total_price=line.total_price,
            paid=line.paid,
            debt=line.debt,
            event_date=None,
            due_date=None,
            notes="To'y",
            status=line.status,
        )
    )
    try:
        client_raw = client_path.read_bytes()
        assert client_raw.startswith(b"%PDF")
        assert b"FAZO CREATIVE GROUP" in client_raw
        assert len(client_raw) > 1000
    finally:
        client_path.unlink(missing_ok=True)


async def main() -> None:
    check_money()
    check_orders_pdf()
    before = await counts()
    await check_empty_excel()
    await check_accounting()
    after = await counts()
    if before != after:
        raise SystemExit(f"ROLLBACK_FAILED before={before} after={after}")
    print("VERIFY_OK")
    print(f"ROWS_UNCHANGED={before}")
    await engine.dispose()


if __name__ == "__main__":
    asyncio.run(main())
