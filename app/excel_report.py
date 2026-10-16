import tempfile
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Never

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.reports import ExcelData, PeriodKind

NAVY = "0E4D6C"
GOLD = "C4A35A"
HEADER = "145A7A"
ZEBRA = "F4F7FB"
TOTAL = "E5F2E9"
WHITE = "FFFFFF"
LINE = "D5DDE5"

THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=WHITE)
SUBTITLE_FONT = Font(name="Calibri", size=11, italic=True, color="E8EEF3")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
CELL_FONT = Font(name="Calibri", size=11, color="1D2430")
TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color="143D28")
TITLE_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FILL = PatternFill("solid", fgColor=HEADER)
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
TOTAL_FILL = PatternFill("solid", fgColor=TOTAL)
GOLD_FILL = PatternFill("solid", fgColor=GOLD)


@dataclass(frozen=True)
class TableSpec:
    title: str
    headers: tuple[str, ...]
    rows: list[list[object]]
    money_columns: frozenset[int]
    date_columns: frozenset[int]
    tab_color: str


def cashflow_note(kind: PeriodKind) -> str:
    match kind:
        case "daily":
            return (
                "Buyurtma qiymati kassaga tushgan pul emas. "
                "Pul oqimi = bugungi to'lovlar − xarajatlar − jamoa haqi."
            )
        case "monthly":
            return (
                "Kelishilgan narx va shu oy tushgan to'lov alohida yuradi. "
                "Jamoa haqi xarajat ichida ikkinchi marta ayirilmaydi."
            )
        case "yearly":
            return (
                "Kelishilgan narx va shu yil tushgan to'lov alohida yuradi. "
                "Jamoa haqi xarajat ichida ikkinchi marta ayirilmaydi."
            )
        case _ as unexpected:
            never: Never = unexpected
            raise RuntimeError(f"Noma'lum hisobot davri: {never}")


def _summary_rows(data: ExcelData) -> list[list[object]]:
    figures = data.figures
    return [
        ["Davr", data.period.label],
        ["Yangi buyurtmalar soni", figures.new_orders],
        ["Buyurtmalar qiymati (kassa emas)", figures.order_value],
        ["Tushgan to'lovlar", figures.received],
        ["Xarajatlar", figures.expenses],
        ["Jamoa haqi", figures.team_paid],
        ["Jami chiqim", figures.outgoing],
        ["Pul oqimi", figures.net],
        ["Joriy qarzdorlik", figures.outstanding_debt],
        ["Qarzdorlar soni", figures.debtor_count],
        ["Izoh", cashflow_note(data.period.kind)],
    ]


def _order_rows(lines) -> list[list[object]]:
    return [
        [
            line.number,
            line.client_name,
            line.phone,
            line.service,
            line.total_price,
            line.paid,
            line.debt,
            line.event_date,
            line.created_on,
            line.status,
            line.notes,
        ]
        for line in lines
    ]


def tables_for(data: ExcelData) -> list[TableSpec]:
    order_headers = (
        "Buyurtma",
        "Mijoz",
        "Telefon",
        "Xizmat",
        "Umumiy narx",
        "Jami to'langan",
        "Qolgan qarz",
        "Tadbir sanasi",
        "Buyurtma sanasi",
        "Holat",
        "Izoh",
    )
    money = frozenset({5, 6, 7})
    dates = frozenset({8, 9})
    return [
        TableSpec(
            "Umumiy hisobot",
            ("Ko'rsatkich", "Qiymat"),
            _summary_rows(data),
            frozenset(),
            frozenset(),
            NAVY,
        ),
        TableSpec(
            "Buyurtmalar",
            order_headers,
            _order_rows(data.orders),
            money,
            dates,
            "1F6F8B",
        ),
        TableSpec(
            "To'lovlar",
            ("ID", "Buyurtma", "Mijoz", "Summa", "Usul", "Izoh", "Sana"),
            [
                [
                    line.payment_id,
                    line.order_number,
                    line.client_name,
                    line.amount,
                    line.method,
                    line.note,
                    line.paid_on,
                ]
                for line in data.payments
            ],
            frozenset({4}),
            frozenset({7}),
            "2E7D4F",
        ),
        TableSpec(
            "Qarzdorlar",
            order_headers,
            _order_rows(data.debtors),
            money,
            dates,
            "A15C38",
        ),
        TableSpec(
            "Xarajatlar",
            ("ID", "Kategoriya", "Summa", "Izoh", "Sana"),
            [
                [
                    line.expense_id,
                    line.category,
                    line.amount,
                    line.description,
                    line.spent_on,
                ]
                for line in data.expenses
            ],
            frozenset({3}),
            frozenset({5}),
            "8A5A00",
        ),
        TableSpec(
            "Jamoa to'lovlari",
            ("ID", "Xodim", "Summa", "Buyurtma yoki xizmat", "Izoh", "Sana"),
            [
                [
                    line.payment_id,
                    line.person_name,
                    line.amount,
                    line.work_for,
                    line.note,
                    line.paid_on,
                ]
                for line in data.team
            ],
            frozenset({3}),
            frozenset({6}),
            "5C4B8A",
        ),
    ]


def _paint(cell, *, font: Font, fill: PatternFill | None = None, align: str = "left") -> None:
    cell.font = font
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = THIN
    if fill is not None:
        cell.fill = fill


def _is_money_value(value: object) -> bool:
    return isinstance(value, Decimal) and not isinstance(value, bool)


def _write_table(sheet: Worksheet, spec: TableSpec, subtitle: str) -> None:
    width = len(spec.headers)
    sheet.sheet_properties.tabColor = spec.tab_color
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    title = sheet.cell(1, 1, "FAZO CREATIVE GROUP")
    _paint(title, font=TITLE_FONT, fill=TITLE_FILL, align="left")
    subtitle_cell = sheet.cell(2, 1, subtitle)
    _paint(subtitle_cell, font=SUBTITLE_FONT, fill=TITLE_FILL, align="left")
    for column in range(1, width + 1):
        sheet.cell(1, column).fill = TITLE_FILL
        sheet.cell(2, column).fill = TITLE_FILL
        sheet.cell(1, column).border = THIN
        sheet.cell(2, column).border = THIN
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 20
    sheet.row_dimensions[3].height = 8
    for column in range(1, width + 1):
        accent = sheet.cell(3, column, "")
        accent.fill = GOLD_FILL

    for index, header in enumerate(spec.headers, start=1):
        cell = sheet.cell(4, index, header)
        _paint(cell, font=HEADER_FONT, fill=HEADER_FILL, align="center")
    sheet.row_dimensions[4].height = 22
    last_data_row = 4 + len(spec.rows)
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(width)}{max(4, last_data_row)}"

    for row_index, row in enumerate(spec.rows):
        excel_row = 5 + row_index
        fill = ZEBRA_FILL if row_index % 2 else None
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(excel_row, column, _excel_value(value))
            align = "right" if column in spec.money_columns or _is_money_value(value) else "left"
            _paint(cell, font=CELL_FONT, fill=fill, align=align)
            if _is_money_value(value):
                cell.number_format = "#,##0.00"
            if column in spec.date_columns and isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"
        sheet.row_dimensions[excel_row].height = 18

    total_row = 5 + len(spec.rows)
    label = sheet.cell(total_row, 1, "Jami")
    _paint(label, font=TOTAL_FONT, fill=TOTAL_FILL)
    for column in range(2, width + 1):
        cell = sheet.cell(total_row, column, None)
        _paint(cell, font=TOTAL_FONT, fill=TOTAL_FILL, align="right")
        if column not in spec.money_columns:
            continue
        if spec.rows:
            letter = get_column_letter(column)
            cell.value = f"=SUBTOTAL(109,{letter}5:{letter}{total_row - 1})"
        else:
            cell.value = Decimal("0.00")
        cell.number_format = "#,##0.00"
    if spec.title == "Umumiy hisobot":
        note = sheet.cell(total_row, 2, "Pul oqimi va chiqim yuqorida alohida hisoblangan")
        _paint(note, font=TOTAL_FONT, fill=TOTAL_FILL)
    sheet.row_dimensions[total_row].height = 20

    for column in range(1, width + 1):
        letter = get_column_letter(column)
        longest = len(spec.headers[column - 1])
        for row in spec.rows:
            if column - 1 < len(row):
                longest = max(longest, len(_display(row[column - 1])))
        minimum = 16 if column in spec.money_columns else 12
        sheet.column_dimensions[letter].width = min(max(longest + 3, minimum), 42)

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_view.showGridLines = False
    sheet.oddHeader.left.text = "FAZO CREATIVE GROUP"
    sheet.print_title_rows = "1:4"


def _excel_value(value: object) -> object:
    return value


def _display(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:,.2f}"
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value)


def build_workbook(data: ExcelData) -> Workbook:
    workbook = Workbook()
    specs = tables_for(data)
    subtitle = f"{data.period.label}. {cashflow_note(data.period.kind)}"
    first = workbook.active
    first.title = specs[0].title
    _write_table(first, specs[0], subtitle)
    for spec in specs[1:]:
        sheet = workbook.create_sheet(spec.title)
        _write_table(sheet, spec, subtitle)
    workbook.properties.creator = "FAZO CREATIVE GROUP BOT"
    workbook.properties.title = f"FAZO CREATIVE GROUP — {data.period.label}"
    workbook.properties.subject = "Moliyaviy hisobot"
    return workbook


def write_excel(data: ExcelData) -> Path:
    handle = tempfile.NamedTemporaryFile(prefix="fazo_hisobot_", suffix=".xlsx", delete=False)
    path = Path(handle.name)
    handle.close()
    try:
        workbook = build_workbook(data)
        workbook.save(path)
        workbook.close()
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return path
